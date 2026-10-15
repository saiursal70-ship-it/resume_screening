import io
import os
from typing import Union
import docx
import PyPDF2


def extract_text_from_pdf(source: Union[bytes, str]) -> str:
    """Extract plain text from a PDF file path or bytes."""
    if isinstance(source, str):
        with open(source, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            pages = [page.extract_text() or "" for page in reader.pages]
    else:
        reader = PyPDF2.PdfReader(io.BytesIO(source))
        pages = [page.extract_text() or "" for page in reader.pages]

    return "\n\n".join(p.strip() for p in pages if p.strip())


def extract_text_from_docx(source: Union[bytes, str]) -> str:
    """Extract plain text from a DOCX file path or bytes."""
    if isinstance(source, str):
        doc = docx.Document(source)
    else:
        doc = docx.Document(io.BytesIO(source))

    paragraphs = []
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                paragraphs.append(" | ".join(row_texts))

    return "\n\n".join(paragraphs)


def extract_text_from_file(filename: str, source: Union[bytes, str]) -> str:
    """Detect file extension and extract plain text."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(source)
    elif ext == ".docx":
        return extract_text_from_docx(source)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only .pdf and .docx are supported.")
